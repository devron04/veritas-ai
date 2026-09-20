"""
Document parser — extracts plain text from PDF, DOCX, and TXT files.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF using pypdf or PyMuPDF (fitz)."""
    text_parts: list[str] = []

    # Try pypdf first (pure python, zero C compilation issues)
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page_num, page in enumerate(reader.pages, 1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
            else:
                logger.warning("Page %d: no text extracted.", page_num)
        full_text = "\n\n".join(text_parts)
        if full_text.strip():
            return full_text
    except ImportError:
        pass

    # Fallback to PyMuPDF if installed
    try:
        import fitz  # PyMuPDF
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            for page_num, page in enumerate(doc, 1):
                page_text = page.get_text("text")
                if page_text.strip():
                    text_parts.append(page_text)
                else:
                    logger.warning("Page %d: no text extracted.", page_num)
        full_text = "\n\n".join(text_parts)
        if full_text.strip():
            return full_text
    except ImportError:
        pass

    full_text = "\n\n".join(text_parts)
    if not full_text.strip():
        raise ValueError(
            "No text could be extracted from the PDF. "
            "The file may contain only images/scanned pages."
        )
    return full_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx.

    Extracts from paragraphs AND tables to ensure no content is missed.
    """
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(file_bytes))
    text_parts: list[str] = []

    # Extract body paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text)

    # Extract text from tables (cells contain paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        text_parts.append(p.text)

    # Extract headers and footers from each section
    for section in doc.sections:
        for header_footer in (section.header, section.footer):
            if header_footer is not None:
                for p in header_footer.paragraphs:
                    if p.text.strip():
                        text_parts.append(p.text)

    full_text = "\n\n".join(text_parts)

    if not full_text.strip():
        raise ValueError("No text could be extracted from the DOCX file.")
    return full_text


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from a plain text file with encoding detection."""
    # Try UTF-8 first, then fall back to latin-1
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return text
        except (UnicodeDecodeError, ValueError):
            continue
    raise ValueError("Could not decode the text file with any supported encoding.")


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Route to the correct parser based on file extension.

    Returns the extracted plain text.
    Raises ValueError if the format is unsupported or extraction fails.
    """
    suffix = Path(filename).suffix.lower()

    parsers = {
        ".pdf": extract_text_from_pdf,
        ".docx": extract_text_from_docx,
        ".txt": extract_text_from_txt,
    }

    parser = parsers.get(suffix)
    if parser is None:
        raise ValueError(
            f"Unsupported file format: '{suffix}'. "
            f"Accepted formats: {', '.join(parsers.keys())}"
        )

    text = parser(file_bytes)
    logger.info("Extracted %d characters from '%s'.", len(text), filename)
    return text
